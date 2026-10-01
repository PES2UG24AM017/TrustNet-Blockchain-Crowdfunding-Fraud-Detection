"""Patch para [96] — shorten XAI limitation note."""
from docx import Document

SRC = "TrustNet_audit.docx"
DST = r"C:\Users\amrit\OneDrive\Desktop\College\CCNCS_Research\TrustNet.docx"

doc = Document(SRC)

NEW_96 = (
    "Attributions are relative to the chosen baseline and may shift with a different "
    "reference point; the model was trained on synthetic data, so attributions reflect "
    "learned patterns rather than confirmed real-world fraud cases."
)

for i, para in enumerate(doc.paragraphs):
    t = para.text.strip()
    if t.startswith("A limitation of the IG analysis"):
        old = t
        for run in para.runs: run.text = ""
        if para.runs: para.runs[0].text = NEW_96
        else: para.add_run(NEW_96)
        print(f"[{i}] OLD({len(old.split())}w): {old[:80]}")
        print(f"[{i}] NEW({len(NEW_96.split())}w): {NEW_96[:80]}")
        break

wc = sum(len(p.text.split()) for p in doc.paragraphs)
doc.save(DST)
print(f"Saved. Total words: {wc}")
